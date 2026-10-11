import io
import re
import hashlib
from pathlib import Path
from datetime import datetime

import pandas as pd
import streamlit as st
import fitz  # PyMuPDF
from docx import Document as DocxDocument

st.set_page_config(page_title="Legal E-Discovery Screening Tool", page_icon="⚖️", layout="wide")

DEFAULT_TERMS = "contract, payment, delivery, breach, agreement, invoice, refund, termination, notice, debt"

st.title("⚖️ Legal E-Discovery & Document Relevance Screening Tool")
st.caption("A student prototype for screening batches of documents for potential relevance. It does not determine legal relevance or admissibility.")

with st.sidebar:
    st.header("Case screening settings")
    matter_name = st.text_input("Matter / case name", "Emeka v. ABC Construction Ltd.")
    legal_issue = st.selectbox(
        "Legal issue",
        ["Breach of contract", "Debt recovery", "Landlord and tenant", "Employment dispute",
         "Property dispute", "Fraud allegation", "Other"]
    )
    raw_terms = st.text_area("Keywords (comma-separated)", DEFAULT_TERMS, height=120)
    threshold = st.slider("Flag threshold (keyword-match score)", 0, 100, 25, 5)
    st.markdown("---")
    st.caption("Supported formats: TXT, PDF, DOCX. Upload only files you are authorised to process.")

def extract_text(uploaded_file):
    name = uploaded_file.name
    data = uploaded_file.getvalue()
    suffix = Path(name).suffix.lower()
    if suffix == ".txt":
        text = data.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        text_parts = []
        with fitz.open(stream=data, filetype="pdf") as pdf:
            for page in pdf:
                text_parts.append(page.get_text())
        text = "\n".join(text_parts)
    elif suffix == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    else:
        raise ValueError("Unsupported file type")
    return text, hashlib.sha256(data).hexdigest(), len(data)

def find_entities(text):
    # Lightweight pattern-based hints, not a full named-entity recognition model.
    dates = re.findall(r'\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{1,2}\s+[A-Za-z]+\s+\d{4}|[A-Za-z]+\s+\d{1,2},?\s+\d{4})\b', text)
    money = re.findall(r'(?:₦|N|NGN|USD|\$|£)\s?[\d,]+(?:\.\d{2})?', text, flags=re.I)
    emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b', text)
    return dates, money, emails

uploads = st.file_uploader(
    "Upload a batch of sample documents",
    type=["txt", "pdf", "docx"],
    accept_multiple_files=True,
    help="For a quick demonstration, upload the sample .txt files included with this project."
)

col1, col2, col3 = st.columns(3)
col1.metric("Selected legal issue", legal_issue)
col2.metric("Keywords configured", len([x for x in raw_terms.split(",") if x.strip()]))
col3.metric("Documents uploaded", len(uploads) if uploads else 0)

if uploads:
    terms = list(dict.fromkeys([t.strip().lower() for t in raw_terms.split(",") if t.strip()]))
    results = []
    detail_rows = []
    for file in uploads:
        try:
            text, sha256, size = extract_text(file)
            lower = text.lower()
            matched = [term for term in terms if term in lower]
            score = round((len(matched) / len(terms)) * 100, 1) if terms else 0
            dates, money, emails = find_entities(text)
            flagged = score >= threshold
            if not text.strip():
                status = "Review needed — no extractable text"
            elif flagged:
                status = "Flagged for human review"
            else:
                status = "Below screening threshold"
            results.append({
                "Document": file.name,
                "Keyword-match score (%)": score,
                "Matched keywords": ", ".join(matched) if matched else "None",
                "Screening result": status,
                "Text length (characters)": len(text),
                "File size (bytes)": size,
                "SHA-256": sha256
            })
            detail_rows.append({
                "name": file.name, "text": text, "matched": matched, "score": score,
                "status": status, "dates": dates, "money": money, "emails": emails, "hash": sha256
            })
        except Exception as exc:
            results.append({
                "Document": file.name,
                "Keyword-match score (%)": 0,
                "Matched keywords": "—",
                "Screening result": f"Processing error: {exc}",
                "Text length (characters)": 0,
                "File size (bytes)": len(file.getvalue()),
                "SHA-256": hashlib.sha256(file.getvalue()).hexdigest()
            })

    df = pd.DataFrame(results)
    st.subheader("Screening results")
    st.info("A high score means more of the selected keywords appeared in the document. It is not a probability of legal relevance, truth, or admissibility.")
    st.dataframe(df.drop(columns=["SHA-256"]), use_container_width=True, hide_index=True)

    flagged_count = sum(1 for x in detail_rows if x["score"] >= threshold)
    m1, m2, m3 = st.columns(3)
    m1.metric("Documents processed", len(results))
    m2.metric("Flagged for review", flagged_count)
    m3.metric("Below threshold", max(0, len(results) - flagged_count))

    st.subheader("Document review")
    for item in detail_rows:
        with st.expander(f'{item["name"]} — {item["status"]} ({item["score"]}%)'):
            st.write("**Matched keywords:**", ", ".join(item["matched"]) if item["matched"] else "None")
            st.write("**Possible dates detected:**", ", ".join(dict.fromkeys(item["dates"])) if item["dates"] else "None detected")
            st.write("**Possible monetary amounts detected:**", ", ".join(dict.fromkeys(item["money"])) if item["money"] else "None detected")
            st.write("**Possible email addresses detected:**", ", ".join(dict.fromkeys(item["emails"])) if item["emails"] else "None detected")
            st.write("**SHA-256 file fingerprint:**")
            st.code(item["hash"])
            preview = item["text"][:1800]
            st.text_area("Extracted text preview (first 1,800 characters)", preview, height=160, key="preview_" + re.sub(r'[^A-Za-z0-9]', '_', item["name"]))

    csv_bytes = df.to_csv(index=False).encode("utf-8")
    st.download_button("Download screening report (CSV)", data=csv_bytes,
                       file_name="ediscovery_screening_report.csv", mime="text/csv")
else:
    st.markdown("### How to test the prototype")
    st.write("1. Upload several sample documents. 2. Choose the legal issue and keywords in the sidebar. 3. Review flagged documents and export the CSV report.")
    st.write("The project folder includes fictional sample documents concerning a hypothetical contract dispute.")

with st.expander("Limitations, privacy, and responsible use"):
    st.markdown("""
- This is an educational prototype, not a production legal-discovery system.
- Keyword matching can miss relevant documents that use different wording and can flag irrelevant documents that happen to contain a keyword.
- Date, money, and email detection uses simple patterns and may miss or misread information.
- Scanned PDFs may need OCR; this version extracts selectable PDF text only.
- A human reviewer must assess context, privilege, confidentiality, authenticity, relevance, and admissibility.
- Do not upload confidential, privileged, personal, or sensitive client information unless you have authority and an approved secure environment.
- The SHA-256 value helps identify changes to a file; it does not establish that the file is authentic or that its contents are true.
""")

st.caption("Educational prototype • Built for CSC 202 project demonstration")
